from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from copy import deepcopy
from hashlib import sha256
from lxml import etree as E
from docx import Document
from docx.shared import Inches
import json

ROOT=Path(__file__).resolve().parents[2]
REF=Path(r'C:\Users\ad999\Downloads\Motorq_Hackathon_Solution_Document_Template.docx')
OUT=ROOT/'deliverables/ZSpeed_Motorq_Solution_Document.docx'
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W=NS['w']; q=lambda s:'{'+W+'}'+s
z=ZipFile(REF); parts={n:z.read(n) for n in z.namelist()}
root=E.fromstring(parts['word/document.xml']); body=root.find('w:body',NS)
ps=body.findall('w:p',NS); ts=body.findall('w:tbl',NS)
sect=deepcopy(body.find('w:sectPr',NS))

def distill():
    inv={n:{'bytes':len(b),'sha256':sha256(b).hexdigest(),'role':'editable' if n in ['word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'] else 'preserve only'} for n,b in parts.items()}
    (ROOT/'deliverables/qa/package_inventory.json').write_text(json.dumps(inv,indent=2))
    (ROOT/'deliverables/qa/artifact.md').write_text(f'''# Template contract
Reference: {REF}
SHA256: {sha256(REF.read_bytes()).hexdigest()}
Reference rendered with Word to reference.pdf and inspected all 10 pages in reference/page-NN.png.
Packaged renderer unavailable because LibreOffice is absent. Word export and bundled Poppler provide visual verification.
One A4 section 11906 by 16838 twips; margins top1440 left1300 right1300 bottom1300. Header and footer distances600. Preserve sectPr exactly.
Calibri11 default, source Heading1 16pt, Heading2 12pt, cover28pt bold and22pt subtitle. Body before60 after200 line300 auto. Cover title before3200 after120 and subtitle after2200. Retain source typography and paragraph properties; use real Title style for cover.
Header two supplied logos preserved. Footer Page field preserved. All noneditable parts remain byte identical, including styles, numbering, header, footer, footnotes, endnotes, theme and images.
Source lists and tables may be cloned, rows expanded, text replaced. Table grid source-derived with expanded cell padding and explicit D9D9D9 borders. No fixed row heights. Repeat headers.
Ordered content: cover, linked static contents, sections1 through17 with original subsections. Keep bookmarks and contents anchors. Replace instructional bullets and answer placeholders with solution-specific prose, comparison rows and working screenshot. Retain unknown team and contact fields and link slots as editable blanks. Add future roadmap within section12.
Slots: document.xml body paragraph0 title,1 subtitle,2-8 metadata;10-42 linked contents;44 onward headings/bookmarks and associated prose/table slots. Every prompt and sample table row is editable; header/footer/section geometry preserve only. Title/screenshot and architecture flow are allowed additions within existing content slots. More pages permitted to accommodate completed answers without reducing type.
Field inventory: footer PAGE only; Word export refreshes field for QA without resaving source. No content controls, body images, rich text fields, or comments.
Fidelity gates: reference hash unchanged, preserve-only package equality, source section geometry, contents anchors, all final pages inspected. Expected difference: instructional content replaced, more table rows, screenshot, completed document pagination.
''',encoding='utf-8')

def para(text,idx=50,style=None):
    p=deepcopy(ps[idx]); pr=p.find('w:pPr',NS)
    for c in list(p):
        if c.tag!=q('pPr'): p.remove(c)
    if pr is None: pr=E.SubElement(p,q('pPr'))
    if style:
        st=pr.find('w:pStyle',NS)
        if st is None: st=E.SubElement(pr,q('pStyle'))
        st.set(q('val'),style)
    r=E.SubElement(p,q('r'))
    if idx in [0,1]:
        source_rpr=ps[idx].find('w:r/w:rPr',NS)
        if source_rpr is not None:r.append(deepcopy(source_rpr))
    t=E.SubElement(r,q('t')); t.text=text
    t.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
    body.append(p); return p

def h(idx): body.append(deepcopy(ps[idx]))
def p(text): return para(text)
def b(text): return para(text,46)
page_calls=0
def page():
    global page_calls
    page_calls+=1
    if page_calls>2:return
    x=para('',9); r=E.SubElement(x,q('r')); E.SubElement(r,q('br')).set(q('type'),'page')
def table(idx,rows,widths=None):
    tbl=deepcopy(ts[idx]); old=tbl.findall('w:tr',NS)
    for r in old: tbl.remove(r)
    pr=tbl.find('w:tblPr',NS)
    borders=pr.find('w:tblBorders',NS)
    if borders is None: borders=E.SubElement(pr,q('tblBorders'))
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=borders.find('w:'+side,NS)
        if e is None:e=E.SubElement(borders,q(side))
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(q(k),v)
    mar=pr.find('w:tblCellMar',NS)
    if mar is None:mar=E.SubElement(pr,q('tblCellMar'))
    for side in ['top','bottom','left','right']:
        e=mar.find('w:'+side,NS)
        if e is None:e=E.SubElement(mar,q(side))
        e.set(q('w'),'85'); e.set(q('type'),'dxa')
    for i,values in enumerate(rows):
        r=deepcopy(old[0 if i==0 else min(1,len(old)-1)])
        rp=r.find('w:trPr',NS)
        if rp is None:rp=E.SubElement(r,q('trPr'))
        for ht in rp.findall('w:trHeight',NS):rp.remove(ht)
        E.SubElement(rp,q('cantSplit'))
        if i==0:E.SubElement(rp,q('tblHeader'))
        for cell,txt in zip(r.findall('w:tc',NS),values):
            cp=cell.find('w:tcPr',NS)
            for child in list(cell):
                if child.tag!=q('tcPr'):cell.remove(child)
            par=E.SubElement(cell,q('p')); pp=E.SubElement(par,q('pPr'))
            E.SubElement(pp,q('spacing'),{q('after'):'40',q('line'):'250',q('lineRule'):'auto'})
            run=E.SubElement(par,q('r')); rpr=E.SubElement(run,q('rPr'))
            E.SubElement(rpr,q('sz'),{q('val'):'18' if idx==2 else '20'})
            if i==0:E.SubElement(rpr,q('b'))
            E.SubElement(run,q('t')).text=txt
            va=cp.find('w:vAlign',NS)
            if va is None:va=E.SubElement(cp,q('vAlign'))
            va.set(q('val'),'center')
        tbl.append(r)
    body.append(tbl)

def picture():
    image=Path(r'C:\Users\ad999\.codex\visualizations\2026\10\01\01a0f793-e457-7b93-bb89-62175fe770c9\repair-units.png')
    d=Document(); pp=d.add_paragraph(); pp.add_run().add_picture(str(image),width=Inches(6.0))
    el=deepcopy(pp._p)
    relroot=E.fromstring(parts['word/_rels/document.xml.rels'])
    rid='rIdZSpeedScreenshot'
    for blip in el.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip'):
        blip.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed',rid)
    E.SubElement(relroot,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',Id=rid,Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/image',Target='media/zspeed-repair.png')
    parts['word/_rels/document.xml.rels']=E.tostring(relroot,xml_declaration=True,encoding='UTF-8',standalone=True)
    parts['word/media/zspeed-repair.png']=image.read_bytes(); body.append(el)

def build():
    for e in list(body):body.remove(e)
    para('ZSpeed Fleet Operations',0,'Title')
    para('Connected Vehicle Intelligence Solution',1)
    for text in ['Submission format: PDF export of this document','Team name: ________________________________________','Team members and roles: _____________________________','Contact email: ______________________________________','Problem space: Fleet telemetry quality and vehicle operations','GitHub repository URL: _______________________________','Demo video URL: ____________________________________','Submission date: ____________________________________']:
        para(text,3)
    page()
    for src in ps[10:43]:body.append(deepcopy(src))
    page()
    h(44)
    p('ZSpeed Fleet Operations gives fleet managers one place to onboard vehicles, connect OEM accounts, inspect vehicle health and trips, and restore readings held after an OEM changes its data format. We built a working fleet interface, platform backend, independent OEM simulator and simulator control console.')
    p('The repair workspace lets a manager match renamed fields, choose measurement units, translate status codes and decide how to handle added or removed readings. A preview checks the proposed changes before a versioned repair is enabled and valid historical records are restored. Original payloads and accepted history remain available.')
    p('Vehicle summaries and trip bundles are built in background jobs and served from a separate read database copy, with optional Redis caching. This removes repeated reconstruction from raw telemetry during page visits. The current implementation uses a durable SQLite queue and a replica on the same host. Kafka and a distributed database are proposed next steps.')
    h(51);h(52)
    p('A fleet manager needs usable vehicle and trip information when OEM payloads change because a renamed field, different unit or missing reading can otherwise interrupt reporting and require technical investigation. Operations staff need to keep valid data visible while the affected readings are reviewed.')
    p('Primary user: fleet manager. Other stakeholders: operations analysts, platform administrators and OEM integration teams. The prototype demonstrates data quality and recovery; fleet-wide financial or safety benefits have not been measured.')
    h(57)
    table(0,[['Evidence or assumption','Method','Finding','Confidence'],['Format changes can block readings','Synthetic OEM changes and repair tests','Renames, units and status translations can be repaired locally','High for tested cases'],['Bad coordinates can break a map','Coordinate validation and map tests','Invalid points are handled without crashing vehicle views','High for tested cases'],['Saved trip bundles reduce read work','Read scaling tests and source review','Vehicle visits read saved summaries instead of rebuilding raw history','High structurally'],['Managers can complete repairs unaided','Guided browser walkthrough','Flow works end to end; independent user study still needed','Provisional']])
    h(62)
    p('Success is measured by usable trip data, recovery outcomes, query cost and visible freshness. No 10,000 or 100,000 vehicle deployment, cost saving, throughput or percentile latency is claimed. A pilot should measure those outcomes against the existing raw-event lookup path.')
    page();h(67);h(68)
    p('A manager imports vehicles, confirms onboarding, connects an OEM account and chooses the readings to receive. The simulator then sends events. Accepted events update vehicle state and trips; held events appear as repair categories with concrete actions.')
    b('Open a vehicle to see saved health details, a trip list and a route with important events and data quality information.')
    b('Open the affected repair category. Match the old field to its new location, select the correct unit or translate the new status value.')
    b('Preview the sampled readings. Resolve any failed checks, then choose Restore valid readings. Track the recovery by its job ID.')
    picture()
    p('Working product example: the unit repair view shows the received speed, the selected unit and the converted value before the manager previews the repair.')
    h(73)
    p('Managers can fix supported format changes inside the platform and recover usable history without writing transformation code. The interface keeps valid vehicle information available and reports what remains held. Conflicting identifiers and malformed records are separated from repairable mapping changes because they cannot safely be solved by choosing another field.')
    h(79)
    p('Three implemented ideas distinguish this demonstrator: repair choices tailored to the failure category with before-and-after samples; versioned recovery that preserves raw payloads and accepted history; and grouped background activity that turns many events for one vehicle into a readable job row. These are implementation choices, not claims of industry-first functionality.')
    page();h(84)
    p('Done means implemented in the current code. Video timestamps below are suggested recording segments, to be confirmed after the demo is recorded. Code paths are relative to the repository; backend services are under backend/src/services/.')
    table(2,[['ID','Feature','User story','Priority','Status','Code path','Video'],['F01','CSV onboarding','Add vehicles with checked rows','Must','Done','import.service.ts','1:00'],['F02','OEM connections','Know which vehicles are receiving data','Must','Done','connection.service.ts','1:20'],['F03','Trip map and quality','Inspect usable routes and missing data','Must','Done','read-model.service.ts','1:45'],['F04','Guided format repair','Fix renamed fields and changed units','Must','Done','MappingRepair.tsx; mapping-repair.service.ts','2:00'],['F05','Historical recovery','Restore valid held readings safely','Must','Done','quarantine.service.ts','3:00'],['F06','Grouped jobs','See progress for one vehicle activity','Should','Done','background-activity.service.ts','3:30'],['F07','Fast saved reads','Load saved trip details in one request','Must','Done','read-model.service.ts; read-cache.service.ts','3:45'],['F08','Docker and AWS package','Start the packaged stack','Should','Partial','docker-compose.yml; deploy-aws.ps1','4:00']])
    p('F08 is packaged with launch scripts and health checks. A live Docker deployment and AWS deployment still require verification in an environment with the required tooling and account access.')
    h(87);h(88)
    p('System context: the fleet manager uses the fleet web application. An operator uses the simulator console to configure fictional OEM sources. The backend connects those sources to fleet-scoped ingestion, repair and read APIs. No live commercial OEM account or external identity provider is integrated in the demonstrator.')
    p('Container flow: Fleet UI → HTTP and Socket.IO → backend API. Simulator UI → HTTP → simulator backend → connector ingestion → operational SQLite. A supervised worker consumes durable jobs and updates a materialized read database → replica SQLite → optional Redis → fleet read API.')
    p('An event is retained, checked for duplication, normalized using its selected mapping and either accepted or quarantined. Accepted changes request coalesced projection work. The UI reads the latest completed snapshot; hop latency has not been benchmarked.')
    page();h(93)
    table(3,[['Layer','Choice','Reason'],['Web interfaces','React, TypeScript, Vite','Shared typed models and separate fleet and simulator screens'],['Platform API','Node.js, Express, Socket.IO','HTTP workflows and fleet-scoped live updates'],['Processing','Supervised worker and durable SQLite jobs','Persisted retries and restart recovery without a broker dependency'],['Data and cache','SQLite WAL and optional Redis','Simple deployment; separate saved read views and bounded cache fallback'],['Mapping','Deterministic rules and versioned configurations','Predictable previews and repeatable validation'],['Deployment','Docker Compose, Caddy, AWS CloudFormation helper','Same stack locally or on a single cloud host']])
    h(95)
    p('The relational core associates fleets with vehicles and OEM connections. Connections own raw events and mapping versions. Accepted measurements feed vehicle state, trips and routes; held records belong to quarantine incidents and recovery jobs. Raw payloads remain the source for replay.')
    p('Three database files separate operational state, materialized read views and the same-host read replica. Saved trip bundles deliberately combine route points, important events and quality statistics so a selected trip needs one detail fetch. Trip lists are paginated. Redis entries expire quickly and fall back to the replica when Redis is unavailable.')
    table(4,[['Query','Before ms','After ms','Change'],['Vehicle detail','Not measured','Not measured','Read saved summary and health statistics'],['Selected trip','Not measured','Not measured','Fetch one prebuilt route and event bundle']])
    p('Illustrative sizing assumption: 10,000 vehicles sending one 1 KB event every 30 seconds produce about 333 events/second and 28.8 GB/day of raw payloads before indexes, replicas and backups. At 100,000 vehicles, those figures become about 3,333 events/second and 288 GB/day. These are capacity estimates, not tested capabilities. Retention and archival policies must be set for a pilot.')
    h(101)
    p('Docker Compose runs backend, simulator, a Caddy container serving both built interfaces, and Redis. Only web ports need to be public; backend and simulator ports stay on the internal network. Health checks establish startup readiness. Persistent volumes retain operational and simulator data.')
    p('On Windows, Start-ZSpeed.cmd or start.ps1 prepares configuration and starts the stack. start.sh provides the Linux entry point. The AWS helper packages the source and provisions a single EC2 host through CloudFormation. A domain and DNS configuration enable HTTPS through Caddy. Docker, cloud smoke tests and backup restoration remain release checks.')
    page();h(105);h(106)
    p('The backend uses routes for request boundaries, services for ingestion, repair and projections, database modules for persistence, and connector adapters for OEM-specific behavior. Some services still use better-sqlite3 directly; a fully database-independent repository layer is future migration work.')
    table(5,[['Layer','Responsibility','Boundary'],['Presentation','Fleet and simulator screens; HTTP and socket handlers','Preserve selected views and check user permissions'],['Application services','Import, connections, ingestion, repair and recovery','Validate changes before publishing a mapping version'],['Processing','Trip building, route checks, projection jobs','Run outside client page requests'],['Persistence and adapters','SQLite, Redis, encrypted secrets and OEM connectors','Keep raw history and scope reads to the fleet']])
    p('Repository layout: frontend/src contains fleet views; simulator-ui/src contains simulator controls; backend/src contains routes, services, database code, simulator code and tests; docker contains deployment configuration. Top-level launch scripts start the packaged environment.')
    h(111)
    p('Idempotency combines event identity with a payload hash. An exact repeat reuses the accepted result; different content with the same identity is held for review. Role and fleet checks enforce least privilege. Environment configuration supplies deployment settings. Dedicated mapping and route functions keep transformation behavior testable.')
    h(115)
    table(6,[['Pattern','Purpose','Code'],['Adapter','Translate OEM-specific formats into canonical readings','connector contract and format mapping'],['Separate write and read models','Keep vehicle browsing away from raw-event reconstruction','read-model.service.ts'],['Durable queue with retries','Resume processing and recovery after interruptions','worker.service.ts'],['Cache bypass and bounded waits','Serve replica data if Redis fails or stalls','read-cache.service.ts']])
    h(117)
    p('Client APIs use JSON and fleet-scoped authorization. Live notifications use Socket.IO. The canonical event retains source identity, event time, vehicle association, supported measurements and quality information. Mapping versions define source fields, conversions and status translations; no Kafka topics or Avro registry are deployed today.')
    p('Normal flow: simulator sends event → backend retains raw payload → identity and format checks → normalized measurements → queued projection → read replica → UI refreshes the affected data. Failure flow: format fails → category-specific repair → sampled preview → immutable version enabled → durable replay → valid records restored, remaining invalid records held.')
    h(122)
    p('Trip building sorts and deduplicates points, filters implausible GPS jumps, segments by ignition or time gaps and calculates distance using the haversine formula. Sorting costs O(n log n); linear checks cost O(n). Route simplification uses Ramer Douglas Peucker, which can be O(n²) in the worst case. Stored bundles avoid repeating this work on every visit. Large-input timing and memory profiles remain to be measured.')
    page();h(129)
    table(7,[['Requirement','Pilot target','Current result','Verification'],['API response','p95 below 200 ms under agreed pilot load','Not benchmarked','Load and soak test required'],['Fresh vehicle view','Show completed snapshots and freshness','Implemented; eventual consistency','Read scaling and browser checks'],['Recovery','Resume jobs after worker interruption','Covered for durable queue and replay','Focused integration tests'],['Cache outage','Read from replica when Redis is unavailable','Bounded fallback implemented','Read cache tests; Docker case pending'],['Availability','99.9 percent service objective','Single host remains a failure point','Multi-host design and failover test required']])
    p('The template’s 100,000 events/second target has not been demonstrated. A repeatable benchmark should record hardware, event mix, fleet size, read/write concurrency, p95 and p99 latency, queue age, replica lag and database growth. Comparison should use the same workload before and after saved read models.')
    h(134)
    p('Authentication and role checks protect fleet APIs and socket rooms. Connection secrets are encrypted in a vault. Production launch scripts create credentials and encryption keys. A domain deployment can use HTTPS; local development and simulated OEM authorization do not establish an OIDC or mTLS deployment.')
    table(0,[['Threat','Control','Remaining work','Scope'],['Cross-fleet access','Fleet-scoped queries and socket rooms','Independent security review','Tenant isolation'],['Unsafe repair','Role checks, previews, immutable mapping versions','Audit review and approval policies for pilots','Mapping changes'],['Duplicate or changed event identity','Payload hashing and quarantine','Provider correction for true identity conflicts','Ingestion'],['Credential disclosure','Encrypted vault and redacted source values','Managed key rotation and secrets service','OEM access'],['Resource exhaustion','Background batches and bounded cache waits','Production load limits and abuse tests','Availability']])
    p('Location history is sensitive. Before a pilot, define retention, deletion, access review, backup retention and data processing responsibilities. The demonstrator uses synthetic telemetry; GDPR or DPDP compliance certification is not claimed. Full database encryption, regional disaster recovery and penetration testing are future requirements.')
    page();h(141)
    table(8,[['Test area','Tools','Checks','Result','CI'],['Repair backend','Node test runner and TypeScript','27','Passed in focused validation','Not verified'],['Repair plans and guidance','Node test runner','19','Passed in focused validation','Not verified'],['Read scaling','Node test runner','22','Passed in earlier validation','Not verified'],['Grouped activity and refresh','Node test runner','8','Passed in earlier validation','Not verified'],['End-to-end repair','Browser and isolated synthetic fleet','1 walkthrough','Preview, replay and resolution completed','Manual'],['Production build','TypeScript and Vite','Backend and frontend','Passed after repair changes','Not verified'],['Cloud, load and security','Planned','No benchmark count','Docker/Redis container check skipped without Docker','Pending']])
    p('The latest repair validation passed 46 focused checks. These counts describe separate validation runs, not a freshly rerun complete suite or a line-coverage percentage. Cases cover renamed nested fields, unit conversions, added and removed readings, status changes, GPS pairs, timestamp requirements, identity conflicts, stale drafts, sensitive fields and tenant permissions.')
    p('Example: a new OEM payload moves speed to metrics.speed and sends 80 km/h. The manager selects that field and unit, sees 80 → 80 km/h, previews all selected sample records and restores valid readings. Invalid samples stay held. The browser walkthrough also verified a newly added heading, translated engine status and a removed odometer field.')
    h(146)
    p('The Background Jobs tab groups continuous processing for the same vehicle under a stable job ID. It shows the purpose, first start, latest activity, status, record counts and recovery progress. Data Pipeline updates preserve the current view and repair form instead of reloading the entire page every two seconds.')
    p('To investigate a slow vehicle view, check the snapshot timestamp, find that vehicle’s job, inspect its status and record totals, then review worker logs and replica progress. A growing queue suggests processing pressure; fresh projections with slow responses suggest read or cache latency. Distributed traces, latency percentiles and alert thresholds are proposed operational additions.')
    h(150)
    p('There is no trained model or LLM in the runtime repair path. Suggestions use deterministic field matching and explicit conversion rules. Ambiguous source choices stay with the manager. Future assisted suggestions would require measured accuracy, clear provenance and approval before any mapping is enabled.')
    page();h(157)
    p('The current decisions favor a deployable single-host demonstrator with recoverable processing and useful read isolation. The roadmap below describes proposed work; these capabilities are not part of the current deployment.')
    table(3,[['Decision','Current choice','Consequence'],['Operational database','SQLite WAL','Simple setup; one host and writer limit'],['Messaging','Durable SQLite queue','Restart recovery without a broker; multi-host scaling requires replacement'],['Read consistency','Asynchronous projections and replica','Pages stay responsive; users may see a recently completed snapshot'],['Repair safety','Previewed immutable mapping versions','Changes remain reviewable and original records are retained'],['Cache role','Optional short-lived Redis entries','Redis failure can bypass to replica; freshness is bounded by both layers']])
    h_future=para('Future product experience',52,'Heading2')
    p('A manager would open a single attention list showing which OEM changed, how many vehicles are affected and which readings need review. Each item would offer a proposed field or unit change with old and new samples, a confidence indicator and a preview of affected trips. The manager would approve a version, schedule recovery and watch one grouped progress row. An audit history would show who changed what and allow a new version to restore the previous rules.')
    table(3,[['Stage','Enhancement','Ready when'],['Pilot readiness','Connect a real OEM; add secure token rotation, audit history, retention and tested backups','An approved real feed passes contract tests and an operator can restore a backup'],['Scale the data layer','Move operational storage to PostgreSQL with an independent read replica; add an outbox and Kafka or a managed queue','Load tests meet defined read/write targets and crash replay preserves idempotency'],['Run across hosts','Deploy API and workers on multiple hosts with load balancing, metrics, alerts and controlled autoscaling','Host failure tests meet recovery targets without losing acknowledged events'],['Improve repair assistance','Offer schema differences and suggested mappings while retaining explicit preview and manager approval','User testing and regression checks show safe, understandable suggestions']])
    p('Proposed cloud flow: OEM webhooks or polling → authenticated ingestion → transactional outbox → Kafka partitioned by vehicle → normalization and projection workers → PostgreSQL primary and independent read replica → Redis → fleet API. Object storage would hold archived raw payloads under a defined retention policy. Stable identifiers, idempotent consumers and dead-letter handling would keep retries safe.')
    p('Kafka versus a simpler managed queue should be decided from ordering, replay, throughput and operating cost needs. PostgreSQL migration must validate counts and hashes, compare old and new read views and support rollback. Availability claims should follow failure testing rather than the choice of a cloud service.')
    page();h(162)
    p('Demo video URL: __________________________________________________')
    p('Alternate video URL: ______________________________________________')
    p('Keep the final recording within five minutes. The outline below is a recording plan, not a completed video or a verified timestamp index.')
    table(9,[['Time','Segment','What to show'],['0:00–0:30','Problem','A fleet manager sees a format change and held readings'],['0:30–1:00','Solution','Fleet view, vehicle health and saved trip route'],['1:00–2:00','Connected data','Import or select vehicles, connect a simulator OEM and show a live event'],['2:00–3:30','Manager repair','Match a renamed field, choose its unit, preview and restore valid readings'],['3:30–4:15','Recovery and architecture','Show the grouped recovery job, restored trip data and current queue/read model'],['4:15–5:00','Evidence and next steps','Show focused test outcomes and the proposed real OEM and distributed deployment']])
    h(166)
    p('GitHub repository URL: _____________________________________________')
    p('Release or submission tag URL: ______________________________________')
    b('Included: service folders, source tests, Docker Compose, local launch scripts and deployment documentation.')
    b('Before submission: reconcile README architecture with the latest read model, record the demo, add the links and confirm the submission tag.')
    b('Before publishing: run the full suite and container smoke test, configure CI, scan dependencies and images, and verify that generated credentials and database files are excluded.')
    h(173)
    p('ZSpeed demonstrates a complete path from connected vehicle ingestion to trip visibility and manager-led recovery after supported OEM format changes. Preserved raw history, versioned repairs and saved read views make the results inspectable and recoverable. The next validation step is a real OEM pilot with measured load, security review and tested backup recovery.')
    page();h(178)
    p('Open-source components include React, Vite, TypeScript, Node.js, Express, better-sqlite3, SQLite, Socket.IO, Leaflet, Redis and Caddy. Exact versions and license obligations should be captured from the lockfiles and deployment images in a submission software bill of materials; no license audit is claimed here.')
    p('AI tools: OpenAI Codex assisted with code changes, test development, troubleshooting and this solution document. The deployed application does not call an LLM to normalize or repair telemetry.')
    p('Demo data: OEM names and telemetry are generated by the simulator. The demonstrated scenarios do not represent live commercial OEM integrations. Confirm that the final recording, repository and any added dataset contain only permitted synthetic or public data.')
    p('Team declaration and approval: ______________________________________')
    h(183)
    p('Implementation references are relative to the repository and describe the current source:')
    b('DEPLOYMENT.md and docker-compose.yml: packaged services, launch configuration and cloud prerequisites.')
    b('READ_SCALING_PLAN.md: saved read views, replica, cache and background processing work.')
    b('frontend/src/components/MappingRepair.tsx and RepairWorkbench.tsx: category-specific manager repair screens.')
    b('backend/src/services/mapping-repair.service.ts and mapping-repair.engine.ts: mapping versions, preview and conversion rules.')
    b('backend/src/services/read-model.service.ts, read-cache.service.ts and worker.service.ts: saved summaries and durable work.')
    b('backend/src/tests/: repair, data quality, tenant isolation, read scaling, queue and trip tests.')
    p('Supporting test report URL: _________________________________________')
    p('Additional documentation URL: ______________________________________')
    body.append(sect)
    parts['word/document.xml']=E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
    with ZipFile(OUT,'w',ZIP_DEFLATED) as zz:
        for name,data in parts.items():zz.writestr(name,data)
    base=ZipFile(REF); final=ZipFile(OUT)
    changed=[n for n in base.namelist() if base.read(n)!=final.read(n)]
    assert set(changed)<= {'word/document.xml','word/_rels/document.xml.rels'},changed
    assert E.tostring(E.fromstring(final.read('word/document.xml')).find('w:body/w:sectPr',NS))==E.tostring(E.fromstring(base.read('word/document.xml')).find('w:body/w:sectPr',NS))
    print('Created',OUT,'Changed parts',changed)

if __name__=='__main__':
    import sys
    if '--distill' in sys.argv:distill()
    else:build()
